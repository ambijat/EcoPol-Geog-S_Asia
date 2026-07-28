from __future__ import annotations

import hashlib
import json
import subprocess
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation


SYMBOL = "<HISTORICAL_RESOURCE_REPOSITORY>/"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def resolve_symbolic(locator: str, repository_root: Path) -> Path:
    if not locator.startswith(SYMBOL):
        raise ValueError("historical source must use the symbolic repository locator")
    relative = Path(locator.removeprefix(SYMBOL))
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("unsafe historical source locator")
    root = repository_root.resolve()
    source = (root / relative).resolve()
    if root not in source.parents:
        raise ValueError("historical source escapes the read-only repository")
    return source


def extract_pptx(path: Path) -> list[dict]:
    presentation = Presentation(path)
    records = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts = []
        images = charts = tables = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if text:
                    texts.append(text)
            if getattr(shape, "shape_type", None) == 13:
                images += 1
            charts += int(bool(getattr(shape, "has_chart", False)))
            tables += int(bool(getattr(shape, "has_table", False)))
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        has_notes = False
        try:
            has_notes = bool(slide.notes_slide.notes_text_frame.text.strip())
        except (AttributeError, ValueError):
            pass
        records.append({
            "number": number, "title": title, "text": "\n".join(texts),
            "layout_name": slide.slide_layout.name or "", "image_count": images,
            "chart_count": charts, "table_count": tables,
            "speaker_note_available": has_notes, "page_status": "TEXT_EXTRACTED",
        })
    return records


def _pdf_page_count(path: Path) -> int:
    result = subprocess.run(["pdfinfo", str(path)], check=True, capture_output=True, text=True)
    for line in result.stdout.splitlines():
        if line.startswith("Pages:"):
            return int(line.split(":", 1)[1].strip())
    raise ValueError("pdfinfo did not report a page count")


def extract_pdf(path: Path, thumbnail_dir: Path | None = None) -> list[dict]:
    records = []
    for number in range(1, _pdf_page_count(path) + 1):
        result = subprocess.run(
            ["pdftotext", "-f", str(number), "-l", str(number), str(path), "-"],
            check=True, capture_output=True, text=True,
        )
        text = result.stdout.strip()
        preview = ""
        if thumbnail_dir is not None:
            thumbnail_dir.mkdir(parents=True, exist_ok=True)
            prefix = thumbnail_dir / f"page-{number:03d}"
            subprocess.run(
                ["pdftoppm", "-f", str(number), "-l", str(number), "-scale-to", "900", "-png", "-singlefile", str(path), str(prefix)],
                check=True, capture_output=True,
            )
            preview = str(prefix.with_suffix(".png"))
        records.append({
            "number": number, "title": text.splitlines()[0] if text else "",
            "text": text, "layout_name": "PDF_PAGE", "image_count": 0,
            "chart_count": 0, "table_count": 0, "speaker_note_available": False,
            "page_status": "TEXT_EXTRACTED" if text else "IMAGE_ONLY — OCR NOT PERFORMED",
            "preview_path": preview,
        })
    return records


def extract_read_only(source: Path, output_dir: Path) -> dict:
    before = sha256(source)
    suffix = source.suffix.lower()
    converted_pdf = None
    if suffix == ".pptx":
        records = extract_pptx(source)
        tool = "python-pptx"
    elif suffix == ".pdf":
        records = extract_pdf(source, output_dir / "thumbnails")
        tool = "pdfinfo+pdftotext+pdftoppm"
    elif suffix == ".odp":
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if not executable:
            raise RuntimeError("LibreOffice is required for read-only ODP inspection")
        output_dir.mkdir(parents=True, exist_ok=True)
        conversion_dir = Path(tempfile.mkdtemp(prefix="odp-render-", dir=output_dir))
        process = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(conversion_dir), str(source)],
            text=True, capture_output=True, timeout=120,
        )
        converted_pdf = conversion_dir / f"{source.stem}.pdf"
        if process.returncode != 0 or not converted_pdf.is_file():
            raise RuntimeError("ODP conversion failed: " + (process.stderr.strip() or process.stdout.strip()))
        records = extract_pdf(converted_pdf, output_dir / "thumbnails")
        tool = "LibreOffice headless → PDF; pdfinfo+pdftotext+pdftoppm"
    else:
        raise ValueError(f"read-only inspection is not implemented for {suffix or 'unknown files'}")
    after = sha256(source)
    if after != before:
        raise RuntimeError("historical source changed during read-only extraction")
    output_dir.mkdir(parents=True, exist_ok=True)
    extracted_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    cache = output_dir / "extract.json"
    cache.write_text(json.dumps({"slides": records}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    result = {"source_sha256": before, "records": records, "tool": tool,
              "timestamp": extracted_at, "cache": cache, "source_format": suffix.removeprefix(".").upper()}
    if converted_pdf is not None:
        result.update(converted_pdf=converted_pdf, converted_pdf_sha256=sha256(converted_pdf),
                      page_count=len(records), conversion_tool="LibreOffice headless")
    return result
