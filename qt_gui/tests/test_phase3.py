from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import unittest
from datetime import date,timedelta
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM","offscreen")

from pptx import Presentation
from PySide6.QtWidgets import QApplication

from gui.services import classify_public_path
from qt_gui.config import AppConfig,STYLE_SHEET,WATERMARK
from qt_gui.database.connection import DatabaseManager
from qt_gui.dialogs.deliverable_preview_dialog import DeliverablePreviewDialog
from qt_gui.main_window import MainWindow
from qt_gui.services.class_record_service import ClassRecordService
from qt_gui.services.deliverable_service import DeliverableService,sha256
from qt_gui.services.historical_deck_service import HistoricalDeckService
from qt_gui.services.ledger_draft_service import LedgerDraftService
from qt_gui.tests.fixtures.phase3_acceptance_fixture import CLASSROOM_WARNING,FIXTURE_LABEL,seed_phase3_fixture


class QtPhaseThreeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        for folder in ("reports","course","course_ledger","resource_registry","config","historical"): (self.root/folder).mkdir(parents=True,exist_ok=True)
        (self.root/"course_ledger/ledger.jsonl").write_text(json.dumps({"block_number":1,"content_hash":"a"*64})+"\n")
        (self.root/"resource_registry/resources.jsonl").write_text("")
        (self.root/"reports/public_include_manifest.txt").write_text("qt_gui/app.py\n")
        (self.root/"reports/public_exclude_manifest.txt").write_text("course/lectures/**/deck_source/*_DRAFT.pptx\ncourse_ledger/drafts/qt/**\n")
        subprocess.run(["git","init","-b","main"],cwd=self.root,check=True,capture_output=True);subprocess.run(["git","config","user.name","Qt Test"],cwd=self.root,check=True);subprocess.run(["git","config","user.email","qt.invalid@example.invalid"],cwd=self.root,check=True);(self.root/"seed.txt").write_text("seed\n");subprocess.run(["git","add","seed.txt"],cwd=self.root,check=True);subprocess.run(["git","commit","-m","seed"],cwd=self.root,check=True,capture_output=True)
        self.db=DatabaseManager(self.root/"local_state/cockpit.sqlite3",self.root);self.db.initialise();self.ids=seed_phase3_fixture(self.db);self.service=DeliverableService(self.db)

    def tearDown(self):self.temp.cleanup()

    def generate(self,status="DRAFT"):return self.service.generate(1,"A",status)

    def test_evidence_chain_gate_and_instructor_original_exception(self):
        gate=self.service.gate(1,"A");self.assertTrue(gate["enabled"]);self.assertEqual(gate["evidence_complete_count"],1)
        with self.db.connection() as c:c.execute("delete from slide_source_links");c.execute("delete from slide_historical_links");c.execute("update slide_plan_entries set resource_ids='[]',historical_slide_sources='[]'");c.commit()
        self.assertFalse(self.service.gate(1,"A")["enabled"])
        with self.db.connection() as c:c.execute("update slide_plan_entries set source_independence_status='INSTRUCTOR_ORIGINAL',instructor_origin_declaration='Explicit synthetic instructor-origin declaration'");c.commit()
        self.assertTrue(self.service.gate(1,"A")["enabled"])

    def test_pptx_generation_sidecar_and_teaching_brief(self):
        result=self.generate();self.assertTrue(result["path"].is_file());self.assertEqual(result["version"],"v0.1");sidecar=json.loads(result["sidecar"].read_text());self.assertEqual(sidecar["generation_gate"],"PASS");self.assertEqual(sidecar["slides"][0]["slide_id"],"P3-S1");self.assertEqual((sidecar["visibility"],sidecar["review_status"],sidecar["classroom_use_status"]),("PRIVATE","REVIEW_REQUIRED","NOT_FOR_CLASSROOM_USE"));brief=result["teaching_brief"].read_text();self.assertIn(FIXTURE_LABEL,brief);self.assertIn("Unresolved verification items",brief)

    def test_superseded_slide_is_preserved_but_excluded_from_generation(self):
        with self.db.connection() as c:
            original=dict(c.execute("select * from slide_plan_entries where slide_id='P3-S1'").fetchone())
            c.execute("update slide_plan_entries set revision_status='SUPERSEDED_BY_STRUCTURED_REVISION' where slide_id='P3-S1'")
            c.execute("""insert into slide_plan_entries(slide_id,lecture_pair_id,part,sequence,title,purpose,action,
            note_ids,resource_ids,historical_slide_sources,visual_type,speaker_note,citation_footer,verification_status,
            approval_status,generation_sequence) values ('P3-S1A',?,'A',2,?,?,'ADD','[\"P3-N1\"]','[\"P3-R1\"]',
            '[\"P3-HS1\"]','CONCEPT',?,?, 'REQUIRES_INSTRUCTOR_REVIEW','WORKING_DRAFT',1)""",
            (self.ids["pair"],FIXTURE_LABEL,CLASSROOM_WARNING,CLASSROOM_WARNING,"Synthetic replacement citation"))
            replacement=c.execute("select id from slide_plan_entries where slide_id='P3-S1A'").fetchone()[0]
            c.execute("insert into slide_source_links(slide_plan_id,resource_id,created_at) values (?,?,?)",(replacement,self.ids["resource"],"test"))
            c.execute("insert into slide_note_links(slide_plan_id,note_id,created_at) values (?,?,?)",(replacement,self.ids["note"],"test"))
            c.execute("insert into slide_historical_links(slide_plan_id,historical_slide_id,created_at) values (?,?,?)",(replacement,self.ids["historical"],"test"));c.commit()
        gate=self.service.gate(1,"A");self.assertTrue(gate["enabled"]);self.assertEqual([p["slide_id"] for p in gate["plans"]],["P3-S1A"])
        result=self.generate("REVISED");sidecar=json.loads(result["sidecar"].read_text());self.assertEqual([s["slide_id"] for s in sidecar["slides"]],["P3-S1A"])
        with self.db.connection() as c:
            preserved=c.execute("select * from slide_plan_entries where slide_id='P3-S1'").fetchone()
            for field in ("slide_id","sequence","title","purpose","action","note_ids","resource_ids","historical_slide_sources","speaker_note","citation_footer"):
                self.assertEqual(preserved[field],original[field])

    def test_deterministic_versioning_and_nonoverwrite(self):
        first=self.generate("DRAFT");before=sha256(first["path"]);second=self.generate("REVISED");self.assertEqual((first["version"],second["version"]),("v0.1","v0.2"));self.assertNotEqual(first["path"],second["path"]);self.assertEqual(sha256(first["path"]),before)

    def test_approved_generation_remains_blocked(self):
        gate=self.service.gate(1,"A","APPROVED");self.assertFalse(gate["enabled"]);self.assertTrue(any("F4" in e for e in gate["errors"]))

    def test_pdf_render_thumbnail_count_and_preview(self):
        bundle=self.generate();rendered=self.service.render_pdf(bundle["bundle_id"]);thumbs=self.service.generate_thumbnails(bundle["bundle_id"]);self.assertEqual(rendered["page_count"],2);self.assertEqual(len(thumbs),2);dialog=DeliverablePreviewDialog(self.root,self.service.bundle(bundle["bundle_id"]));self.assertEqual(len(dialog.paths),2);dialog.close()

    def test_pdf_render_failure_retains_pptx_and_records_warning(self):
        bundle=self.generate();before=sha256(bundle["path"])
        with patch("qt_gui.services.deliverable_service.convert_pdf",side_effect=RuntimeError("synthetic render failure")):
            with self.assertRaises(RuntimeError):self.service.render_pdf(bundle["bundle_id"])
        self.assertEqual(sha256(bundle["path"]),before);self.assertIn("synthetic render failure",self.service.bundle(bundle["bundle_id"])["rendering_warnings"])

    def test_validation_and_approval_blocking_and_override(self):
        bundle=self.generate();report=self.service.validate(bundle["bundle_id"]);self.assertEqual(report["status"],"REVIEW_REQUIRED");self.assertTrue(any(f["code"]=="ABSENT_REFERENCES_SLIDE" for f in report["findings"]))
        with self.assertRaises(ValueError):self.service.record_deck_decision(bundle["bundle_id"],"INSTRUCTOR_APPROVED","")
        self.service.record_deck_decision(bundle["bundle_id"],"INSTRUCTOR_APPROVED","Synthetic fixture warning override only")
        self.assertEqual(self.service.bundle(bundle["bundle_id"])["deck_decision"],"INSTRUCTOR_APPROVED")

    def test_validation_rules_and_checksum_error(self):
        bundle=self.generate()
        with self.db.connection() as c:
            c.execute("delete from slide_note_links");c.execute("delete from slide_source_links");c.execute("""update slide_plan_entries set purpose='',citation_footer='',verification_status='NOT_YET_VERIFIED',
            action='RETURN_TO_NOTES',visual_type='MAP',visual_asset_path='',title='EMPTY PLACEHOLDER',speaker_note=?""",("x"*1300,));c.commit()
        bundle["path"].write_bytes(bundle["path"].read_bytes()+b"checksum-change")
        report=self.service.validate(bundle["bundle_id"]);codes={f["code"] for f in report["findings"]}
        self.assertTrue({"MISSING_PURPOSE","MISSING_NOTES","MISSING_EVIDENCE","MISSING_CITATION","UNRESOLVED_FACTUAL_CLAIM","UNVERIFIED_SLIDE","EXCESSIVE_TEXT","FONT_SIZE_RISK","MISSING_MAP","EMPTY_PLACEHOLDER","ABSENT_REFERENCES_SLIDE","OUTPUT_CHECKSUM_MISMATCH"}<=codes)
        self.assertEqual(report["status"],"FAIL")
        with self.assertRaises(ValueError):self.service.record_deck_decision(bundle["bundle_id"],"INSTRUCTOR_APPROVED","cannot override errors")

    def test_revision_and_slide_review_workflow(self):
        bundle=self.generate();self.service.record_slide_decision(bundle["bundle_id"],self.ids["slide"],"RETURN_TO_NOTES","Synthetic revision");self.service.record_deck_decision(bundle["bundle_id"],"RETURNED_FOR_REVISION","Synthetic return");detail=self.service.bundle(bundle["bundle_id"]);self.assertEqual(detail["deck_decision"],"RETURNED_FOR_REVISION");self.assertEqual(detail["slide_decisions"][0]["decision"],"RETURN_TO_NOTES")

    def _make_odp(self) -> Path:
        pptx=self.root/"historical/fixture.pptx";presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text=FIXTURE_LABEL;slide.placeholders[1].text=CLASSROOM_WARNING;presentation.save(pptx)
        executable=shutil.which("libreoffice") or shutil.which("soffice");self.assertIsNotNone(executable)
        result=subprocess.run([executable,"--headless","--convert-to","odp","--outdir",str(pptx.parent),str(pptx)],capture_output=True,text=True,timeout=120);self.assertEqual(result.returncode,0,result.stderr);return pptx.with_suffix(".odp")

    def test_odp_conversion_provenance_and_source_immutability(self):
        odp=self._make_odp();before=sha256(odp)
        with self.db.connection() as c:
            pair=c.execute("select id from lecture_pairs where lecture_number=1").fetchone()[0];c.execute("""insert into historical_decks(deck_id,lecture_pair_id,title,source_locator,file_type,source_layer,content_sha256)
            values ('ODP-D1',?,?, '<HISTORICAL_RESOURCE_REPOSITORY>/fixture.odp','.odp','SYNTHETIC',?)""",(pair,FIXTURE_LABEL,before));deck=c.execute("select id from historical_decks where deck_id='ODP-D1'").fetchone()[0];c.commit()
        service=HistoricalDeckService(self.db,self.root/"historical",self.root/"qt_gui/local_state/historical_derivatives");result=service.inspect(deck);self.assertEqual(result["source_format"],"ODP");self.assertEqual(sha256(odp),before)
        with self.db.connection() as c:record=c.execute("select * from odp_conversion_records where deck_id=?",(deck,)).fetchone();self.assertEqual(record["derivative_status"],"READ_ONLY_DERIVATIVE");self.assertEqual(record["source_sha256"],before);self.assertNotIn(str(self.root),record["derivative_pdf"])

    def test_class_record_date_safeguard(self):
        service=ClassRecordService(self.db);future=(date.today()+timedelta(days=2)).isoformat()
        with self.assertRaises(ValueError):service.save(1,"A",{"scheduled_date":future,"actual_date":future,"instructor_validation_status":"INSTRUCTOR_VALIDATED"})
        service.save(1,"A",{"scheduled_date":future,"actual_date":future,"instructor_validation_status":"INSTRUCTOR_VALIDATED","retrospective_entry_explanation":"Synthetic safeguard override"});self.assertEqual(service.load(1,"A")["instructor_validation_status"],"INSTRUCTOR_VALIDATED")

    def test_draft_ledger_isolation_and_schema(self):
        ClassRecordService(self.db).save(1,"A",{"scheduled_date":date.today().isoformat(),"actual_date":date.today().isoformat(),"instructor_validation_status":"INSTRUCTOR_VALIDATED"});canonical=(self.root/"course_ledger/ledger.jsonl").read_bytes();path=LedgerDraftService(self.db).prepare(1,"CLASS_SESSION_COMPLETED");payload=json.loads(path.read_text());self.assertIn("course_ledger/drafts/qt",path.as_posix());self.assertEqual(payload["approval_status"],"DRAFT");self.assertIsNone(payload["canonical_hash"]);self.assertEqual((self.root/"course_ledger/ledger.jsonl").read_bytes(),canonical)

    def test_symbolic_and_public_private_classification(self):
        include={"qt_gui/services/deliverable_service.py"};exclude={"course/lectures/**/deck_source/*_DRAFT.pptx","course_ledger/drafts/qt/**"};self.assertEqual(classify_public_path("qt_gui/services/deliverable_service.py",include,exclude),"PUBLIC_CANDIDATE");self.assertEqual(classify_public_path("course/lectures/lecture_01/part_a_tuesday/deck_source/x_DRAFT.pptx",include,exclude),"PRIVATE_LOCAL")

    def test_transaction_rollback(self):
        with self.assertRaises(RuntimeError):
            with self.db.connection() as c:c.execute("update lecture_pairs set weekly_title='ROLLBACK' where lecture_number=1");raise RuntimeError("cancel")
        with self.db.connection() as c:self.assertNotEqual(c.execute("select weekly_title from lecture_pairs where lecture_number=1").fetchone()[0],"ROLLBACK")

    def test_qt_deliverables_and_class_record_interactions(self):
        config=AppConfig(self.root,self.db.database_path,WATERMARK,STYLE_SHEET,self.root/"historical",self.root/"qt_gui/local_state/historical_derivatives");window=MainWindow(config);window.show();self.app.processEvents();window.open_workspace("Deliverables",1);self.assertEqual(window.navigation.currentRow(),8);self.assertTrue(window.deliverables_view.generate_button.isEnabled());window.open_workspace("Class Record",1);self.assertEqual(window.navigation.currentRow(),9);self.assertEqual(window.class_record_view.status.currentText(),"NOT_REVIEWED");window.close()


if __name__=="__main__":unittest.main()
