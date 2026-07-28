from __future__ import annotations

import hashlib
import json
import os
import subprocess
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pptx import Presentation
from PySide6.QtCore import Qt
from PySide6.QtGui import QImage, QPainter, QPdfWriter
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from gui.services import add_historical_slide
from qt_gui.config import AppConfig, STYLE_SHEET, WATERMARK
from qt_gui.database.connection import DatabaseManager
from qt_gui.main_window import MainWindow
from qt_gui.models.table_model import DictTableModel
from qt_gui.services.historical_deck_service import HistoricalDeckService
from qt_gui.services.historical_extraction import extract_read_only, resolve_symbolic
from qt_gui.services.note_service import NoteService
from qt_gui.services.resource_service import ResourceService
from qt_gui.services.slide_plan_service import SlidePlanService


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class QtPhaseTwoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        (self.root / "reports").mkdir(); (self.root / "course").mkdir()
        (self.root / "course_ledger").mkdir(); (self.root / "resource_registry").mkdir()
        (self.root / "course_ledger/ledger.jsonl").write_text(json.dumps({"block_number":1,"content_hash":"a"*64})+"\n")
        (self.root / "resource_registry/resources.jsonl").write_text("")
        self.db = DatabaseManager(self.root / "local_state/cockpit.sqlite3", self.root); self.db.initialise()
        self.resources = ResourceService(self.db); self.notes = NoteService(self.db); self.slides = SlidePlanService(self.db)
        self.repository = self.root / "historical"; self.repository.mkdir()
        self.derivatives = self.root / "qt_gui/local_state/historical_derivatives"
        subprocess.run(["git","init","-b","main"],cwd=self.root,check=True,capture_output=True)
        subprocess.run(["git","config","user.name","Qt Test"],cwd=self.root,check=True)
        subprocess.run(["git","config","user.email","qt.invalid@example.invalid"],cwd=self.root,check=True)
        (self.root/"seed.txt").write_text("seed\n"); subprocess.run(["git","add","seed.txt"],cwd=self.root,check=True); subprocess.run(["git","commit","-m","seed"],cwd=self.root,check=True,capture_output=True)

    def tearDown(self): self.temp.cleanup()

    def seed_resource_deck(self, source: Path | None = None, file_type=".pdf"):
        source = source or self.repository / "synthetic.pdf"
        if not source.exists(): self.make_pdf(source, text=True)
        sha = digest(source)
        with self.db.connection() as c:
            pair=c.execute("select id from lecture_pairs where lecture_number=1").fetchone()[0]
            c.execute("""insert into resources(resource_id,title,source_locator,provenance,source_layer,file_type,historical_status,content_sha256)
            values ('R1','Synthetic resource','<HISTORICAL_RESOURCE_REPOSITORY>/synthetic.pdf','Synthetic fixture','SYNTHETIC',?,'HISTORICAL_READ_ONLY',?)""",(file_type,sha))
            resource=c.execute("select id from resources where resource_id='R1'").fetchone()[0]
            c.execute("insert into resource_assignments(resource_id,lecture_pair_id,part) values (?,?,'UNCLASSIFIED')",(resource,pair))
            c.execute("""insert into historical_decks(deck_id,lecture_pair_id,title,source_locator,file_type,source_layer,content_sha256)
            values ('D1',?,'Synthetic deck','<HISTORICAL_RESOURCE_REPOSITORY>/synthetic.pdf',?,'SYNTHETIC',?)""",(pair,file_type,sha))
            deck=c.execute("select id from historical_decks where deck_id='D1'").fetchone()[0]; c.commit()
        return resource,deck,source

    @staticmethod
    def make_pdf(path: Path, text: bool):
        writer=QPdfWriter(str(path)); painter=QPainter(writer)
        if text:painter.drawText(100,200,"Synthetic South Asia page")
        else:
            image=QImage(80,80,QImage.Format.Format_RGB32); image.fill(Qt.GlobalColor.blue); painter.drawImage(100,100,image)
        painter.end()

    def test_resource_assignment_and_history(self):
        resource,_,_=self.seed_resource_deck(); self.resources.assign(1,resource,"A","MAP","CENTRAL","map evidence")
        row=self.resources.list_for_lecture(1)[0]; self.assertEqual((row["assignment"],row["classification"]),("A","MAP")); self.assertEqual(len(self.resources.detail(resource)["history"]),1)

    def test_invalid_assignment_rolls_back(self):
        resource,_,_=self.seed_resource_deck()
        with self.assertRaises(ValueError):self.resources.assign(1,resource,"C","MAP","CENTRAL")
        self.assertEqual(self.resources.list_for_lecture(1)[0]["assignment"],"UNCLASSIFIED")

    def test_symbolic_path_enforcement(self):
        source=self.repository/"ok.txt"; source.write_text("ok")
        self.assertEqual(resolve_symbolic("<HISTORICAL_RESOURCE_REPOSITORY>/ok.txt",self.repository),source.resolve())
        with self.assertRaises(ValueError):resolve_symbolic(str(source),self.repository)
        with self.assertRaises(ValueError):resolve_symbolic("<HISTORICAL_RESOURCE_REPOSITORY>/../escape",self.repository)

    def test_historical_pptx_metadata_extraction_and_immutability(self):
        path=self.repository/"synthetic.pptx"; deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text="Synthetic title"; slide.placeholders[1].text="Readable body"; deck.save(path)
        before=digest(path); result=extract_read_only(path,self.derivatives/"pptx")
        self.assertEqual(result["records"][0]["title"],"Synthetic title"); self.assertIn("Readable body",result["records"][0]["text"]); self.assertEqual(digest(path),before)

    def test_historical_pdf_metadata_extraction(self):
        path=self.repository/"text.pdf"; self.make_pdf(path,True); result=extract_read_only(path,self.derivatives/"pdf")
        self.assertEqual(len(result["records"]),1); self.assertEqual(result["records"][0]["page_status"],"TEXT_EXTRACTED"); self.assertTrue(Path(result["records"][0]["preview_path"]).is_file())

    def test_image_only_pdf_status(self):
        path=self.repository/"image.pdf"; self.make_pdf(path,False); result=extract_read_only(path,self.derivatives/"image")
        self.assertEqual(result["records"][0]["page_status"],"IMAGE_ONLY — OCR NOT PERFORMED")

    def test_derivative_provenance_and_source_immutability(self):
        _,deck,source=self.seed_resource_deck(); before=digest(source); service=HistoricalDeckService(self.db,self.repository,self.derivatives); service.inspect(deck)
        with self.db.connection() as c:
            record=c.execute("select * from historical_derivatives order by id limit 1").fetchone(); self.assertEqual(record["derivative_status"],"READ_ONLY_DERIVATIVE"); self.assertTrue(record["source_locator"].startswith("<HISTORICAL_RESOURCE_REPOSITORY>")); self.assertNotIn(str(self.repository),record["relative_path"])
        self.assertEqual(digest(source),before)

    def test_historical_slide_classification(self):
        resource,deck,_=self.seed_resource_deck()
        with self.db.connection() as c: slide_id=add_historical_slide(c,deck,{"historical_slide_number":1,"title":"Synthetic"})
        service=HistoricalDeckService(self.db,self.repository,self.derivatives)
        with self.db.connection() as c: pk=c.execute("select id from historical_slides where historical_slide_id=?",(slide_id,)).fetchone()[0]
        service.classify_slide(pk,{"current_status":"REVISE","target_part":"A","reason":"Update","linked_resource_ids":["R1"],"confidence":"HIGH"})
        self.assertEqual(service.slides(deck)[0]["current_status"],"REVISE")

    def test_note_creation_source_and_slide_linking(self):
        resource,deck,_=self.seed_resource_deck()
        with self.db.connection() as c: sid=add_historical_slide(c,deck,{"historical_slide_number":1,"title":"Evidence"})
        with self.db.connection() as c: spk=c.execute("select id from historical_slides where historical_slide_id=?",(sid,)).fetchone()[0]
        note_id=self.notes.create(1,{"part":"A","topic":"Region","claim":"Instructor claim","teaching_function":"DEFINITION","created_by":"Instructor","note_origin":"INSTRUCTOR_AUTHORED"})
        with self.db.connection() as c: npk=c.execute("select id from revised_notes where note_id=?",(note_id,)).fetchone()[0]
        self.notes.save(npk,{"part":"A","topic":"Region","claim":"Instructor claim","teaching_function":"DEFINITION","note_origin":"INSTRUCTOR_AUTHORED","verification_status":"VERIFIED","instructor_status":"APPROVED_WORKING_NOTE","resource_pks":[resource],"historical_slide_pks":[spk]})
        detail=self.notes.detail(npk); self.assertEqual(len(detail["resources"]),1); self.assertEqual(len(detail["historical_slides"]),1)

    def test_note_verification_requires_provenance(self):
        note_id=self.notes.create(1,{"part":"A","topic":"No source","claim":"Claim","teaching_function":"CONCEPT","created_by":"Instructor","note_origin":"INSTRUCTOR_AUTHORED"})
        with self.db.connection() as c: pk=c.execute("select id from revised_notes where note_id=?",(note_id,)).fetchone()[0]
        with self.assertRaises(ValueError):self.notes.save(pk,{"part":"A","topic":"No source","claim":"Claim","teaching_function":"CONCEPT","note_origin":"INSTRUCTOR_AUTHORED","verification_status":"VERIFIED","resource_pks":[],"historical_slide_pks":[]})

    def test_slide_creation_ordering_and_evidence_links(self):
        resource,deck,_=self.seed_resource_deck(); note=self.notes.create(1,{"part":"A","topic":"Region","claim":"Claim","teaching_function":"CONCEPT","created_by":"Instructor","note_origin":"INSTRUCTOR_AUTHORED"})
        with self.db.connection() as c: hist=add_historical_slide(c,deck,{"historical_slide_number":1,"title":"Old"}); npk=c.execute("select id from revised_notes where note_id=?",(note,)).fetchone()[0]; hpk=c.execute("select id from historical_slides where historical_slide_id=?",(hist,)).fetchone()[0]
        ids=[]
        for seq in (1,2):ids.append(self.slides.create(1,{"part":"A","sequence":str(seq),"title":f"Slide {seq}","purpose":"Test","action":"ADD","visual_type":"CONCEPT"}))
        with self.db.connection() as c: pks=[c.execute("select id from slide_plan_entries where slide_id=?",(sid,)).fetchone()[0] for sid in ids]
        self.slides.save(pks[0],{"part":"A","title":"Slide 1","purpose":"Test","action":"REVISE","visual_type":"CONCEPT","resource_pks":[resource],"note_pks":[npk],"historical_slide_pks":[hpk]})
        detail=self.slides.detail(pks[0]); self.assertEqual((len(detail["resources"]),len(detail["notes"]),len(detail["historical_slides"])),(1,1,1))
        self.slides.move(pks[1],-1); self.assertEqual(self.slides.list_for_part(1,"A")[0]["title"],"Slide 2")

    def test_phase2_seed_data_from_census(self):
        root=self.root/"seeded"; (root/"reports").mkdir(parents=True); records=[]
        for i in range(6):records.append({"relative_context":f"LECTURE1/{i}/deck.pdf","extension":".pdf","filename":f"deck{i}.pdf","source_layer":"SYNTHETIC","content_sha256":str(i)*64})
        (root/"reports/source_census.json").write_text(json.dumps({"records":records}))
        (root/"reports/historical_lecture_title_extraction.json").write_text(json.dumps({"records":[]}))
        db=DatabaseManager(root/"state/db.sqlite3",root); db.initialise()
        with db.connection() as c:
            self.assertEqual(c.execute("select count(*) from historical_decks").fetchone()[0],6); self.assertEqual(c.execute("select count(*) from revised_notes where note_id like 'IS529N-L01-A-N%'").fetchone()[0],6); self.assertEqual(c.execute("select count(*) from slide_plan_entries where slide_id like 'IS529N-L01-A-S%'").fetchone()[0],15)

    def test_database_transaction_rollback(self):
        with self.assertRaises(RuntimeError):
            with self.db.connection() as c:c.execute("update lecture_pairs set weekly_title='ROLLBACK' where lecture_number=1"); raise RuntimeError("cancel")
        with self.db.connection() as c:self.assertNotEqual(c.execute("select weekly_title from lecture_pairs where lecture_number=1").fetchone()[0],"ROLLBACK")

    def test_qt_model_updates(self):
        model=DictTableModel((("title","Title"),),[{"title":"Old"}]); self.assertEqual(model.data(model.index(0,0)),"Old"); model.replace([{"title":"New"},{"title":"Second"}]); self.assertEqual(model.rowCount(),2); self.assertEqual(model.data(model.index(0,0)),"New")

    def test_cross_workspace_navigation_and_widget_smoke(self):
        config=AppConfig(self.root,self.db.database_path,WATERMARK,STYLE_SHEET,self.repository,self.derivatives); window=MainWindow(config); window.show(); self.app.processEvents()
        self.assertNotEqual(window.navigation.item(3).text(),"Phase 2 placeholder"); window.open_workspace("Resources",1); self.assertEqual(window.navigation.currentRow(),3); window.open_workspace("Historical Decks",1); self.assertEqual(window.navigation.currentRow(),4); window.open_workspace("Revised Notes",1); self.assertEqual(window.navigation.currentRow(),5); window.open_workspace("Slide Plan",1); self.assertEqual(window.navigation.currentRow(),6); QTest.keyClicks(window.slide_plan_view.search,"region"); self.assertEqual(window.slide_plan_view.search.text(),"region"); window.close()


if __name__ == "__main__": unittest.main()
