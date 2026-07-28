from __future__ import annotations

import json
import re
import shutil
import sqlite3
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from gui.services import validate_relative_path


SUPPORTED_VISUAL_TYPES = {
    "TITLE", "SECTION", "CONCEPT", "TEXT_AND_IMAGE", "MAP", "DATA_TABLE",
    "CHART", "CHART_PLACEHOLDER", "COMPARISON", "DISCUSSION", "SYNTHESIS", "REFERENCES",
}

NAVY = RGBColor(22, 43, 66)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
GOLD = RGBColor(217, 119, 6)
PALE_BLUE = RGBColor(231, 241, 255)
PALE_TEAL = RGBColor(226, 247, 243)
PALE_GOLD = RGBColor(255, 245, 222)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(71, 85, 105)


def next_deck_filename(directory: Path, lecture_number: int, part: str, status: str) -> tuple[str, str]:
    status = status.upper()
    if status not in {"DRAFT", "REVISED", "APPROVED"}:
        raise ValueError("invalid deck status")
    prefix = f"IS529N_L{lecture_number:02d}{part}_v"
    versions: list[tuple[int, int]] = []
    pattern = re.compile(rf"^{re.escape(prefix)}(\d+)\.(\d+)_")
    for path in directory.glob(f"{prefix}*.pptx"):
        match = pattern.match(path.name)
        if match:
            versions.append((int(match.group(1)), int(match.group(2))))
    if status == "APPROVED":
        approved = sorted((major, minor) for major, minor in versions if major >= 1)
        version = (1, 0) if not approved else (approved[-1][0], approved[-1][1] + 1)
    else:
        minor = max([minor for major, minor in versions if major == 0], default=0) + 1
        version = (0, minor)
    version_text = f"v{version[0]}.{version[1]}"
    return f"{prefix}{version[0]}.{version[1]}_{status}.pptx", version_text


def _add_footer(slide: Any, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.05), Inches(12.3), Inches(0.25))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text or "SOURCE REQUIRED"
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(70, 86, 100)
    paragraph.alignment = PP_ALIGN.RIGHT


def _shape_text(slide: Any, x: float, y: float, w: float, h: float, text: str, *,
                fill: RGBColor = PALE_BLUE, line: RGBColor = BLUE, size: int = 16,
                bold: bool = False, color: RGBColor = NAVY) -> Any:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    paragraph = frame.paragraphs[0]; paragraph.text = text; paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(size); paragraph.font.bold = bold; paragraph.font.color.rgb = color
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.05)
    return shape


def _visual_title(slide: Any, title: str, subtitle: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.85))
    frame = box.text_frame; frame.clear()
    p = frame.paragraphs[0]; p.text = title; p.font.size = Pt(27); p.font.bold = True; p.font.color.rgb = NAVY
    p = frame.add_paragraph(); p.text = subtitle; p.font.size = Pt(13); p.font.color.rgb = MUTED


def _add_speaker_notes(slide: Any, cue: str) -> None:
    if cue.strip():
        slide.notes_slide.notes_text_frame.text = cue


def _render_region_framework_slide(presentation: Presentation, plan: sqlite3.Row) -> Any | None:
    slide_id = plan["slide_id"]
    if slide_id not in {"IS529N-L01-A-S002A", "IS529N-L01-A-S002B", "IS529N-L01-A-S002C"}:
        return None
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if slide_id.endswith("S002A"):
        _visual_title(slide, plan["title"], "A brief diagnostic warm-up: five dimensions for assessing a region")
        stages = (
            "1\nGeographical\ncoherence",
            "2\nFunctional\ninterdependence",
            "3\nShared\nidentity",
            "4\nInstitutionalisation",
            "5\nRegional\nactorness",
        )
        for index, label in enumerate(stages):
            x = 0.55 + index * 2.52
            _shape_text(slide, x, 2.05, 2.18, 2.35, label, fill=PALE_BLUE if index < 4 else PALE_GOLD,
                        line=BLUE if index < 4 else GOLD, size=15, bold=True)
            if index < 4:
                connector = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, Inches(x + 2.18), Inches(3.22), Inches(x + 2.48), Inches(3.22)
                )
                connector.line.color.rgb = MUTED; connector.line.width = Pt(2)
        _shape_text(slide, 2.4, 5.05, 8.55, 0.7,
                    "Hettne: analytical progression—not an inevitable pathway.",
                    fill=PALE_TEAL, line=TEAL, size=17, bold=True)
    elif slide_id.endswith("S002B"):
        _visual_title(slide, plan["title"], "A brief system lens—not an extended theory segment")
        _shape_text(slide, 0.85, 1.65, 5.55, 2.15,
                    "FOUNDATIONS\n\nSpace  •  Function  •  Scale  •  Territory",
                    fill=PALE_GOLD, line=GOLD, size=19, bold=True)
        _shape_text(slide, 6.9, 1.65, 5.55, 2.15,
                    "REGIONAL SYSTEM\n\nStructure  •  Programme  •  Actor  •  Environment",
                    fill=PALE_TEAL, line=TEAL, size=18, bold=True)
        _shape_text(slide, 1.45, 4.45, 10.45, 0.95,
                    "ACTION SPACE  ↔  ACTION UNIT     |     SHARED IDENTITY  ≠  COLLECTIVE ACTORNESS",
                    fill=PALE_BLUE, line=BLUE, size=18, bold=True)
        _shape_text(slide, 2.15, 5.78, 9.0, 0.55,
                    "From Territorial Space to Regional Actorness",
                    fill=WHITE, line=MUTED, size=15)
    else:
        _visual_title(slide, plan["title"], "Organising gateway: subsequent slides test eight senses of regional formation")
        dimensions = (
            "Geographical\ncontiguity", "Ecological\ninterdependence",
            "Historical and civilisational\nconnectivity", "Economic\ninteraction",
            "Regional\nidentity", "Institutionalisation",
            "Political\nfragmentation", "Collective\nactorness",
        )
        for index, label in enumerate(dimensions):
            row, column = divmod(index, 4)
            fill, line = (PALE_TEAL, TEAL) if index < 6 else (PALE_GOLD, GOLD)
            _shape_text(slide, 0.75 + column * 3.08, 1.55 + row * 1.55, 2.65, 1.05,
                        label, fill=fill, line=line, size=16, bold=True)
        _shape_text(slide, 1.35, 5.15, 10.65, 0.8,
                    "WORKING PROPOSITION TO TEST: strong social-spatial formation; weak institutional and political actorness.",
                    fill=PALE_BLUE, line=BLUE, size=16, bold=True)
    _add_footer(slide, plan["citation_footer"])
    _add_speaker_notes(slide, plan["speaker_note"])
    return slide


def _body_text(plan: sqlite3.Row) -> str:
    notes = json.loads(plan["note_ids"] or "[]")
    resources = json.loads(plan["resource_ids"] or "[]")
    lines = [f"Purpose: {plan['purpose']}"]
    if notes:
        lines.append("Notes: " + ", ".join(notes))
    if resources:
        lines.append("Resources: " + ", ".join(resources))
    if plan["visual_type"] in {"TEXT_AND_IMAGE", "MAP"}:
        lines.append("Visual: " + (plan["visual_asset_path"] or "IMAGE REQUIRED"))
    elif plan["visual_type"] == "DATA_TABLE":
        lines.append("DATA TABLE PLACEHOLDER")
    elif plan["visual_type"] == "CHART_PLACEHOLDER":
        lines.append("CHART PLACEHOLDER")
    return "\n".join(lines)


def generate_pptx(
    connection: sqlite3.Connection,
    project_root: Path,
    lecture_number: int,
    part: str,
    status: str = "DRAFT",
    output_directory: Path | None = None,
) -> dict[str, Any]:
    part = part.upper()
    pair = connection.execute(
        "SELECT * FROM lecture_pairs WHERE lecture_number=?", (lecture_number,)
    ).fetchone()
    part_row = connection.execute(
        "SELECT * FROM lecture_parts WHERE lecture_pair_id=? AND part=?", (pair["id"], part)
    ).fetchone()
    if status.upper() == "APPROVED" and not (
        part_row["fidelity_status"] == "F4"
        and part_row["approval_status"] == "INSTRUCTOR_APPROVED"
    ):
        raise ValueError("approved deck generation requires F4 instructor-approved part status")
    plans = connection.execute(
        """SELECT * FROM slide_plan_entries WHERE lecture_pair_id=? AND part=? AND action!='DELETE'
        AND revision_status!='SUPERSEDED_BY_STRUCTURED_REVISION'
        ORDER BY CASE WHEN generation_sequence>0 THEN generation_sequence ELSE sequence END""",
        (pair["id"], part),
    ).fetchall()
    output_directory = output_directory or (
        project_root / f"course/lectures/lecture_{lecture_number:02d}/"
        f"part_{'a_tuesday' if part == 'A' else 'b_friday'}/deck_source"
    )
    output_directory.mkdir(parents=True, exist_ok=True)
    filename, version = next_deck_filename(output_directory, lecture_number, part, status)
    output_path = output_directory / filename
    if output_path.exists():
        raise FileExistsError(f"refusing to overwrite {output_path}")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = part_row["title"]
    title_slide.placeholders[1].text = (
        f"Economic and Political Geography of South Asia\n{pair['weekly_title']}\n"
        f"{part_row['scheduled_day'].title()} · {part_row['scheduled_time']}\n"
        f"{part_row['identifier']} · {version} {status.upper()}"
    )
    _add_footer(title_slide, "IS529N · instructor review required")

    slide_metadata = []
    for plan in plans:
        visual_type = plan["visual_type"]
        if visual_type not in SUPPORTED_VISUAL_TYPES:
            raise ValueError(f"unsupported visual type: {visual_type}")
        slide = _render_region_framework_slide(presentation, plan)
        if slide is None:
            layout_index = 2 if visual_type == "SECTION" else 1
            slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])
            slide.shapes.title.text = f"[{plan['slide_id']}] {plan['title']}"
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = _body_text(plan)
            _add_footer(slide, plan["citation_footer"])
            _add_speaker_notes(slide, plan["speaker_note"])
        effective_sequence = plan["generation_sequence"] or plan["sequence"]
        slide_metadata.append(
            {
                "slide_id": plan["slide_id"],
                "lecture_id": pair["lecture_id"],
                "part": part,
                "sequence": effective_sequence,
                "purpose": plan["purpose"],
                "visual_type": visual_type,
                "note_ids": json.loads(plan["note_ids"] or "[]"),
                "resource_ids": json.loads(plan["resource_ids"] or "[]"),
                "historical_slide_ids": json.loads(plan["historical_slide_sources"] or "[]"),
                "citation_footer": plan["citation_footer"],
                "verification_status": plan["verification_status"],
                "approval_status": plan["approval_status"],
            }
        )
    presentation.save(output_path)

    relative = output_path.resolve().relative_to(project_root.resolve()).as_posix()
    validate_relative_path(relative)
    sidecar = output_path.with_suffix(".metadata.json")
    brief = output_path.with_suffix(".teaching_brief.md")
    metadata = {
        "course_title": "Economic and Political Geography of South Asia",
        "course_cockpit": "IS529N Semester 2026 Paired Lecture Production Cockpit",
        "lecture_id": pair["lecture_id"], "part": part, "version": version,
        "status": status.upper(), "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "editable_pptx": relative, "slides": slide_metadata,
        "canonical_source_files_modified": False,
    }
    sidecar.write_text(json.dumps(metadata, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    brief_lines = [
        "# Economic and Political Geography of South Asia", "",
        f"## Teaching Brief — {part_row['identifier']}", "",
        f"- Lecture title: {pair['weekly_title']}", f"- Part title: {part_row['title']}",
        f"- Central question: {pair['weekly_central_question'] or 'Pending instructor entry'}",
        f"- Central argument: {pair['weekly_argument'] or 'Pending instructor entry'}", "",
        "## Slide sequence and speaker cues", "",
    ]
    for plan in plans:
        effective_sequence = plan["generation_sequence"] or plan["sequence"]
        resources = ", ".join(json.loads(plan["resource_ids"] or "[]")) or "No linked resource"
        brief_lines.extend([f"### {effective_sequence}. {plan['slide_id']} — {plan['title']}", "",
                            f"- Purpose: {plan['purpose']}", f"- Speaker cue: {plan['speaker_note'] or 'Pending instructor entry'}",
                            f"- Evidence sources: {resources}", f"- Transition: Pending instructor entry",
                            f"- Optional/time-shortening status: {'Optional' if plan['action']=='DELETE' else 'Retain unless instructor shortens'}", ""])
    brief_lines.extend([
        "## Discussion questions", "", "Pending instructor entry.", "",
        "## Bridge to the next session", "", pair["relationship_between_parts"] or "Pending instructor entry.", "",
        "## Unresolved verification items", "",
    ])
    unresolved = [plan["slide_id"] for plan in plans if plan["verification_status"] != "VERIFIED"]
    brief_lines.append(", ".join(unresolved) if unresolved else "None recorded.")
    brief_lines.append("")
    brief.write_text("\n".join(brief_lines), encoding="utf-8")
    sidecar_relative = sidecar.resolve().relative_to(project_root.resolve()).as_posix()
    connection.execute(
        """INSERT INTO deliverables(lecture_pair_id,part,deliverable_type,version,status,
        relative_path,sidecar_path,created_at,approved) VALUES (?,?,?,?,?,?,?,?,?)""",
        (pair["id"], part, "EDITABLE_PPTX", version, status.upper(), relative,
         sidecar_relative, metadata["generated_at"], int(status.upper() == "APPROVED")),
    )
    connection.execute(
        "UPDATE lecture_parts SET lifecycle_status='DECK_DRAFTED' WHERE lecture_pair_id=? AND part=?",
        (pair["id"], part),
    )
    connection.commit()
    return {"path": output_path, "sidecar": sidecar, "teaching_brief": brief, "version": version}


def libreoffice_available() -> bool:
    return bool(shutil.which("libreoffice") or shutil.which("soffice"))


def thumbnail_rendering_available() -> bool:
    return libreoffice_available() and bool(shutil.which("pdftoppm"))


def convert_pdf(pptx_path: Path, output_directory: Path) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice conversion is unavailable")
    target = output_directory / (pptx_path.stem + ".pdf")
    if target.exists():
        raise FileExistsError(f"refusing to overwrite {target}")
    output_directory.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_directory), str(pptx_path)],
        text=True, capture_output=True, timeout=120,
    )
    if result.returncode != 0 or not target.is_file():
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.strip()}")
    return target


def render_pdf_thumbnails(pdf_path: Path, output_directory: Path) -> list[Path]:
    executable = shutil.which("pdftoppm")
    if not executable:
        raise RuntimeError("PDF thumbnail rendering is unavailable")
    output_directory.mkdir(parents=True, exist_ok=True)
    prefix = output_directory / "slide"
    result = subprocess.run(
        [executable, "-png", "-r", "120", str(pdf_path), str(prefix)],
        text=True, capture_output=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(f"thumbnail rendering failed: {result.stderr.strip()}")
    return sorted(output_directory.glob("slide-*.png"))


def convert_and_record_pdf(
    connection: sqlite3.Connection,
    project_root: Path,
    lecture_number: int,
    part: str,
    pptx_result: dict[str, Any],
    status: str,
) -> dict[str, Any]:
    pair = connection.execute(
        "SELECT id FROM lecture_pairs WHERE lecture_number=?", (lecture_number,)
    ).fetchone()
    part_dir = "part_a_tuesday" if part == "A" else "part_b_friday"
    pdf_dir = project_root / f"course/lectures/lecture_{lecture_number:02d}/{part_dir}/classroom_pdf"
    pdf = convert_pdf(pptx_result["path"], pdf_dir)
    preview_dir = project_root / f"course/lectures/lecture_{lecture_number:02d}/{part_dir}/previews/{pdf.stem}"
    thumbnails = render_pdf_thumbnails(pdf, preview_dir) if shutil.which("pdftoppm") else []
    relative = pdf.resolve().relative_to(project_root.resolve()).as_posix()
    connection.execute(
        """INSERT INTO deliverables(lecture_pair_id,part,deliverable_type,version,status,
        relative_path,created_at,approved) VALUES (?,?,?,?,?,?,?,?)""",
        (pair[0], part, "CLASSROOM_PDF", pptx_result["version"], status.upper(), relative,
         datetime.now(timezone.utc).isoformat(timespec="seconds"), int(status.upper() == "APPROVED")),
    )
    connection.commit()
    return {"pdf": pdf, "thumbnails": thumbnails}


def validate_deck_plan(connection: sqlite3.Connection, pair_id: int, part: str) -> dict[str, Any]:
    plans = connection.execute(
        """SELECT * FROM slide_plan_entries WHERE lecture_pair_id=? AND part=? AND action!='DELETE'
        AND revision_status!='SUPERSEDED_BY_STRUCTURED_REVISION'
        ORDER BY CASE WHEN generation_sequence>0 THEN generation_sequence ELSE sequence END""",
        (pair_id, part),
    ).fetchall()
    warnings: list[dict[str, str]] = []
    for plan in plans:
        checks = {
            "slide without purpose": not plan["purpose"].strip(),
            "slide without linked notes": not json.loads(plan["note_ids"] or "[]"),
            "slide without source provenance": not json.loads(plan["resource_ids"] or "[]"),
            "missing citation": not plan["citation_footer"].strip(),
            "unresolved factual claim": plan["verification_status"] != "VERIFIED",
            "missing image": plan["visual_type"] in {"TEXT_AND_IMAGE", "MAP"} and not plan["visual_asset_path"],
            "excessive text": len(plan["speaker_note"]) > 1200 or len(plan["purpose"]) > 350,
            "empty placeholder": "PLACEHOLDER" in plan["title"].upper(),
        }
        for message, failed in checks.items():
            if failed:
                warnings.append({"slide_id": plan["slide_id"], "warning": message})
    return {
        "part": part, "planned_slides": len(plans), "warnings": warnings,
        "status": "PASS" if plans and not warnings else "REVIEW_REQUIRED",
        "layout_validation_scope": "HEURISTIC_ONLY_NO_RENDERED_OVERFLOW_GEOMETRY",
        "font_size_floor": 18,
        "low_resolution_check": "AVAILABLE_ONLY_WHEN_A_RASTER_ASSET_IS_LINKED",
    }
