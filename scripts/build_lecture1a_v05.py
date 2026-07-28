#!/usr/bin/env python3
"""Build the review-only Lecture 1A v0.5 synoptic companion."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DECK_DIR = ROOT / "course/lectures/lecture_01/part_a_tuesday/deck_source"
PDF_DIR = ROOT / "course/lectures/lecture_01/part_a_tuesday/classroom_pdf"
STEM = "IS529N_L01A_v0.5_REVISED"
PPTX = DECK_DIR / f"{STEM}.pptx"
METADATA = DECK_DIR / f"{STEM}.metadata.json"
BRIEF = DECK_DIR / f"{STEM}.teaching_brief.md"
VALIDATION = DECK_DIR / f"{STEM}.validation.json"

IVORY = RGBColor(247, 243, 234)
PAPER = RGBColor(255, 253, 248)
INK = RGBColor(31, 39, 45)
MUTED = RGBColor(91, 96, 94)
BURGUNDY = RGBColor(121, 48, 48)
FOREST = RGBColor(61, 86, 69)
GOLD = RGBColor(174, 135, 69)
PALE_GREEN = RGBColor(231, 237, 229)
PALE_RED = RGBColor(242, 229, 225)
PALE_GOLD = RGBColor(243, 236, 217)
LINE = RGBColor(193, 185, 169)

TITLE_FONT = "DejaVu Serif"
BODY_FONT = "Noto Sans"


SLIDES = [
    {
        "title": "South Asia as a Region",
        "subtitle": "Economic and Political Geography of South Asia",
        "kind": "title",
        "purpose": "Open the bounded Lecture 1A enquiry.",
        "cue": "Frame the lecture as an enquiry into the meanings and limits of region. Do not present regional unity as a settled conclusion.",
        "source": "Canonical F0 pilot master · IS529N-RES-0001",
        "verification": "REVIEW_REQUIRED",
    },
    {
        "title": "The question and the route",
        "subtitle": "In what senses is South Asia a region?",
        "kind": "route",
        "purpose": "Give students one question and three stages of enquiry.",
        "cue": "Use the three stages as the lecture map: locate the space, trace its connections, then test institutions and power.",
        "source": "Lecture 1A v0.4 reconstruction · instructor review evidence",
        "verification": "REVIEW_REQUIRED",
    },
    {
        "title": "A region is more than a location",
        "subtitle": "Theory is a diagnostic tool—not the answer",
        "kind": "lens",
        "purpose": "Condense the theoretical warm-up into one slide.",
        "cue": "Keep this under four minutes. Foundations describe regional space; capacities ask whether interaction develops identity, institutions and collective action.",
        "source": "Schmitt-Egner (2002), pp. 179–187, 195",
        "verification": "CONCEPTUALLY_SUPPORTED",
    },
    {
        "title": "South Asia has more than one boundary",
        "subtitle": "Different questions produce different regional frames",
        "kind": "boundaries",
        "purpose": "Distinguish four legitimate but non-identical definitions.",
        "cue": "Ask which boundary is being used each time the phrase South Asia appears. No single frame should be treated as self-evident.",
        "source": "Inherited master, slides 2–4, 17–19, 47–48 · Van Langenhove (2013)",
        "verification": "REQUIRES_SOURCE_REVIEW",
    },
    {
        "title": "Physical geography connects and divides",
        "subtitle": "Mountains, plains, rivers and seas shape interaction",
        "kind": "physical",
        "purpose": "Present physical geography as structure, not proof of unity.",
        "cue": "Read each feature in two directions: as barrier and corridor. Avoid claiming that contiguity automatically produces political coherence.",
        "source": "Inherited master, slides 19, 22, 37, 43, 48 · evidence partial",
        "verification": "PARTIALLY_SUPPORTED",
    },
    {
        "title": "Interdependence crosses borders",
        "subtitle": "Shared systems create cooperation problems as well as opportunities",
        "kind": "interdependence",
        "purpose": "Show ecological and functional interdependence without overclaiming climatic unity.",
        "cue": "Use rivers, coastal systems, energy and mobility as examples of cross-border dependence. Do not teach monsoon unity as verified.",
        "source": "Bandyopadhyay et al. (2017), introduction and Part III",
        "verification": "PARTIALLY_SUPPORTED",
    },
    {
        "title": "Long histories of movement",
        "subtitle": "Routes and encounters preceded modern borders",
        "kind": "timeline",
        "purpose": "Retain the inherited historical spine in a compact form.",
        "cue": "Treat the sequence as a research frame. The inherited deck represents these themes strongly, but detailed historical claims still require source-by-source verification.",
        "source": "Inherited master, slides 3–24 · dedicated historical synthesis required",
        "verification": "EVIDENCE_REQUIRED",
    },
    {
        "title": "Colonial rule and Partition changed the frame",
        "subtitle": "Consolidation and fragmentation belong to the same regional history",
        "kind": "transition",
        "purpose": "Connect territorial consolidation to later political fragmentation.",
        "cue": "Do not imply that colonial rule created every regional connection. Emphasise reorganisation: administration and infrastructure consolidated space; Partition recast it through sovereign borders.",
        "source": "Inherited master, slides 39, 50, 54–55 · scholarly verification required",
        "verification": "EVIDENCE_REQUIRED",
    },
    {
        "title": "Connected region, divided politics",
        "subtitle": "Interaction persists under rivalry, borders and uneven trust",
        "kind": "tension",
        "purpose": "Present political fragmentation as a counterweight to regional connection.",
        "cue": "Use this as a tension, not a verdict. Current political examples require separate verification before classroom use.",
        "source": "Inherited master, slides 49–56 · current evidence required",
        "verification": "EVIDENCE_REQUIRED",
    },
    {
        "title": "Regional disparities are relational",
        "subtitle": "Comparison matters more than inherited figures",
        "kind": "disparities",
        "purpose": "Preserve comparison categories while excluding obsolete numbers.",
        "cue": "Do not quote the inherited charts as current data. Use population, economy and human development as questions for a later verified data update.",
        "source": "Human Development in South Asia (2015) · current data update required",
        "verification": "UPDATE_REQUIRED",
    },
    {
        "title": "Institutions show ambition—and limits",
        "subtitle": "Formal regionalism does not guarantee effective cooperation",
        "kind": "institutions",
        "purpose": "Distinguish institutional existence from institutional effectiveness.",
        "cue": "Name SAARC and SAFTA as institutional forms, then ask what implementation, trade, transit and political constraints reveal. Update all contemporary assessments.",
        "source": "ADB (2009); Ahmed, Kelegama & Ghani (2010); Bandyopadhyay et al. (2017)",
        "verification": "UPDATE_REQUIRED",
    },
    {
        "title": "One region, uneven coherence",
        "subtitle": "The dimensions do not advance together",
        "kind": "matrix",
        "purpose": "Offer a qualified synthesis without a false score or ranking.",
        "cue": "Read across the five dimensions. Stronger geographical framing does not prove shared identity, institutional depth or collective actorness.",
        "source": "Schmitt-Egner (2002) · bounded Lecture 1A reconstruction",
        "verification": "PROVISIONAL_SYNTHESIS",
    },
    {
        "title": "Evidence discipline",
        "subtitle": "Separate inherited representation from verified teaching claims",
        "kind": "evidence",
        "purpose": "Make the current verification boundary visible.",
        "cue": "Explain that preservation and verification are different acts. The master is accepted as a baseline; its claims are not automatically accepted.",
        "source": "Lecture 1A instructor review record · 2026-07-28",
        "verification": "GOVERNANCE_CONFIRMED",
    },
    {
        "title": "Working conclusion",
        "subtitle": "South Asia is a layered and contested regional formation",
        "kind": "conclusion",
        "purpose": "Close with a qualified proposition and a discussion question.",
        "cue": "State the proposition as a working conclusion, then invite challenge: which dimension is strongest, and what evidence would alter the judgement?",
        "source": "Working proposition · requires instructor and factual review",
        "verification": "REVIEW_REQUIRED",
    },
    {
        "title": "Bridge to Friday · Selected references",
        "subtitle": "From regional formation to the Himalayas",
        "kind": "references",
        "purpose": "Connect Part A to Part B and identify the small source base used here.",
        "cue": "Use the Himalayas as the bridge: boundary, corridor, ecological system and political space. References are selected, not a complete bibliography.",
        "source": "Synoptic companion v0.5 · not for classroom use",
        "verification": "REVIEW_REQUIRED",
    },
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_run(run, size, color=INK, bold=False, font=BODY_FONT, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def text_box(slide, x, y, w, h, text, *, size=20, color=INK, bold=False,
             font=BODY_FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    set_run(run, size, color, bold, font)
    return shape


def box(slide, x, y, w, h, text="", *, fill=PAPER, line=LINE, size=18,
        color=INK, bold=False, radius=True, align=PP_ALIGN.CENTER):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.1)
    if text:
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Inches(0.12)
        frame.margin_top = frame.margin_bottom = Inches(0.06)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        set_run(run, size, color, bold)
    return shape


def rule(slide, x1, y1, x2, y2, color=LINE, width=1.3):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def base_slide(prs, number, title, subtitle, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = IVORY
    rule(slide, 0.58, 0.28, 0.58, 1.23, BURGUNDY, 4)
    text_box(slide, 0.84, 0.24, 11.8, 0.55, title, size=27, bold=True, font=TITLE_FONT)
    text_box(slide, 0.86, 0.83, 11.4, 0.3, subtitle, size=11.5, color=MUTED)
    rule(slide, 0.6, 6.92, 12.73, 6.92, LINE, 0.8)
    text_box(slide, 0.62, 6.98, 9.7, 0.22, source, size=7.7, color=MUTED)
    text_box(slide, 10.25, 6.98, 2.45, 0.22, f"{number:02d}  ·  v0.5 SYNOPSIS", size=7.7,
             color=BURGUNDY, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def add_notes(slide, item):
    slide.notes_slide.notes_text_frame.text = (
        f"Purpose: {item['purpose']}\n\nSpeaker cue: {item['cue']}\n\n"
        f"Evidence: {item['source']}\n\nVerification: {item['verification']}\n"
        "Governance: Synoptic companion to the 65-slide master. Not approved for classroom use."
    )


def render(slide, item):
    kind = item["kind"]
    if kind == "title":
        box(slide, 0.86, 1.55, 0.12, 3.7, fill=BURGUNDY, line=BURGUNDY, radius=False)
        text_box(slide, 1.35, 1.72, 10.4, 1.1, item["title"], size=40, bold=True, font=TITLE_FONT)
        text_box(slide, 1.4, 2.95, 9.6, 0.55, item["subtitle"], size=20, color=FOREST)
        text_box(slide, 1.4, 3.72, 8.8, 0.38, "Lecture 1 · Part A · Tuesday", size=14, color=MUTED)
        text_box(slide, 1.4, 4.38, 9.4, 0.65, "One question: in what senses is South Asia a region?", size=22,
                 color=BURGUNDY, bold=True, font=TITLE_FONT)
        text_box(slide, 1.4, 5.45, 10.0, 0.34, "SYNOPTIC COMPANION · NOT FOR CLASSROOM USE", size=10,
                 color=MUTED, bold=True)
        return
    if kind == "route":
        labels = [("1", "SPACE", "Where is the region?"), ("2", "CONNECTIONS", "What crosses borders?"),
                  ("3", "POWER", "What can act collectively?")]
        for i, (n, head, body) in enumerate(labels):
            x = 0.82 + i * 4.13
            box(slide, x, 1.65, 3.65, 3.1, fill=PAPER, line=[FOREST, GOLD, BURGUNDY][i])
            text_box(slide, x + 0.25, 1.92, 0.6, 0.55, n, size=28, color=[FOREST, GOLD, BURGUNDY][i],
                     bold=True, font=TITLE_FONT)
            text_box(slide, x + 0.25, 2.62, 3.0, 0.45, head, size=14, color=MUTED, bold=True)
            text_box(slide, x + 0.25, 3.25, 3.05, 0.72, body, size=21, bold=True, font=TITLE_FONT)
        box(slide, 2.0, 5.35, 9.35, 0.72, "REGION ≠ UNIFORMITY ≠ POLITICAL UNITY", fill=PALE_GOLD,
            line=GOLD, size=18, color=BURGUNDY, bold=True)
    elif kind == "lens":
        box(slide, 0.85, 1.55, 5.65, 3.55, fill=PALE_GOLD, line=GOLD)
        text_box(slide, 1.25, 1.82, 4.85, 0.4, "FOUNDATIONS", size=14,
                 color=BURGUNDY, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, 1.25, 2.42, 4.85, 1.55, "Space\nFunction\nScale\nTerritory", size=22,
                 bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 6.82, 1.55, 5.65, 3.55, fill=PALE_GREEN, line=FOREST)
        text_box(slide, 7.22, 1.82, 4.85, 0.4, "CAPACITIES", size=14,
                 color=FOREST, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, 7.22, 2.42, 4.85, 1.55, "Interdependence\nIdentity\nInstitutions\nActorness",
                 size=22, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 1.6, 5.52, 10.15, 0.66, "A regional space may exist without becoming a collective actor.",
            fill=PAPER, line=BURGUNDY, size=17, color=BURGUNDY, bold=True)
    elif kind == "boundaries":
        items = [("STATE", "member states and formal institutions"), ("PHYSICAL", "mountains, basins, plains and seas"),
                 ("HISTORICAL", "routes, empires, ideas and mobility"), ("FUNCTIONAL", "trade, labour, water, energy and transit")]
        for i, (head, body) in enumerate(items):
            x, y = 0.9 + (i % 2) * 6.15, 1.55 + (i // 2) * 2.15
            box(slide, x, y, 5.55, 1.72, fill=PAPER, line=[BURGUNDY, FOREST, GOLD, INK][i])
            text_box(slide, x + 0.25, y + 0.2, 1.55, 0.38, head, size=13, bold=True,
                     color=[BURGUNDY, FOREST, GOLD, INK][i])
            text_box(slide, x + 0.25, y + 0.76, 4.95, 0.6, body, size=18, font=TITLE_FONT)
        text_box(slide, 2.0, 6.12, 9.3, 0.38, "Ask first: which boundary—and for what purpose?", size=20,
                 color=BURGUNDY, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
    elif kind == "physical":
        bands = [("HIMALAYAS", "barrier · corridor · watershed", PALE_GOLD, GOLD),
                 ("RIVER PLAINS", "settlement · agriculture · mobility", PALE_GREEN, FOREST),
                 ("INDIAN OCEAN", "coasts · exchange · strategic space", PALE_RED, BURGUNDY)]
        for i, (head, body, fill, color) in enumerate(bands):
            y = 1.55 + i * 1.43
            box(slide, 1.2, y, 10.9, 1.03, fill=fill, line=color)
            text_box(slide, 1.55, y + 0.15, 2.3, 0.32, head, size=13, color=color, bold=True)
            text_box(slide, 4.0, y + 0.12, 7.6, 0.48, body, size=20, bold=True, font=TITLE_FONT)
        box(slide, 2.2, 5.98, 8.95, 0.52, "Physical contiguity structures interaction; it does not settle politics.",
            fill=PAPER, line=LINE, size=15, color=MUTED, bold=True)
    elif kind == "interdependence":
        items = [("RIVERS", "upstream ↔ downstream"), ("COASTS & ENERGY", "shared systems ↔ competing uses"),
                 ("MOBILITY", "routes ↔ borders")]
        for i, (head, body) in enumerate(items):
            x = 0.85 + i * 4.14
            box(slide, x, 1.72, 3.65, 2.75, fill=[PALE_GREEN, PALE_GOLD, PALE_RED][i],
                line=[FOREST, GOLD, BURGUNDY][i])
            text_box(slide, x + 0.25, 2.05, 3.15, 0.4, head, size=14, color=[FOREST, GOLD, BURGUNDY][i],
                     bold=True, align=PP_ALIGN.CENTER)
            text_box(slide, x + 0.3, 2.88, 3.05, 0.78, body, size=20, bold=True, font=TITLE_FONT,
                     align=PP_ALIGN.CENTER)
        box(slide, 2.1, 5.23, 9.1, 0.83, "INTERDEPENDENCE CREATES A NEED TO COOPERATE—NOT A GUARANTEE OF COOPERATION.",
            fill=PAPER, line=BURGUNDY, size=16, color=BURGUNDY, bold=True)
    elif kind == "timeline":
        stages = [("1", "TEXTS & ROUTES"), ("2", "PILGRIMAGE & TRADE"), ("3", "EMPIRES & CITIES"),
                  ("4", "COLONIAL NETWORKS")]
        rule(slide, 1.35, 3.22, 11.95, 3.22, GOLD, 2.5)
        for i, (n, label) in enumerate(stages):
            x = 1.15 + i * 3.12
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(2.72), Inches(0.98), Inches(0.98))
            shape.fill.solid(); shape.fill.fore_color.rgb = PAPER
            shape.line.color.rgb = BURGUNDY; shape.line.width = Pt(2)
            text_box(slide, x, 2.72, 0.98, 0.98, n, size=20, color=BURGUNDY, bold=True,
                     font=TITLE_FONT, align=PP_ALIGN.CENTER)
            text_box(slide, x - 0.72, 4.08, 2.45, 0.68, label, size=14, bold=True,
                     align=PP_ALIGN.CENTER)
        box(slide, 2.55, 5.55, 8.3, 0.62, "Strong inherited representation · detailed claims still need sources",
            fill=PALE_GOLD, line=GOLD, size=15, color=MUTED, bold=True)
    elif kind == "transition":
        box(slide, 0.95, 1.55, 5.35, 3.85, fill=PALE_GREEN, line=FOREST)
        text_box(slide, 1.3, 1.92, 4.65, 0.46, "COLONIAL CONSOLIDATION", size=16, color=FOREST,
                 bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, 1.4, 2.85, 4.45, 1.35, "administration\ninfrastructure\nterritorial ordering",
                 size=21, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 7.0, 1.55, 5.35, 3.85, fill=PALE_RED, line=BURGUNDY)
        text_box(slide, 7.35, 1.92, 4.65, 0.46, "PARTITION & SOVEREIGNTY", size=16, color=BURGUNDY,
                 bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, 7.45, 2.85, 4.45, 1.35, "new borders\nstate rivalry\nreworked mobility",
                 size=21, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 2.0, 5.8, 9.35, 0.58, "Older connections persisted, but their political form changed.",
            fill=PAPER, line=GOLD, size=17, color=INK, bold=True)
    elif kind == "tension":
        box(slide, 0.9, 1.62, 5.35, 3.9, fill=PALE_GREEN, line=FOREST)
        text_box(slide, 1.28, 1.95, 4.6, 0.45, "CONNECTION", size=16, color=FOREST, bold=True,
                 align=PP_ALIGN.CENTER)
        text_box(slide, 1.4, 2.8, 4.35, 1.5, "ecological systems\nsocial networks\ntrade and mobility",
                 size=22, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 7.05, 1.62, 5.35, 3.9, fill=PALE_RED, line=BURGUNDY)
        text_box(slide, 7.43, 1.95, 4.6, 0.45, "FRAGMENTATION", size=16, color=BURGUNDY, bold=True,
                 align=PP_ALIGN.CENTER)
        text_box(slide, 7.55, 2.8, 4.35, 1.5, "sovereign borders\ninterstate rivalry\nuneven trust",
                 size=22, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        text_box(slide, 4.7, 5.76, 3.9, 0.48, "BOTH ARE REGIONAL FACTS", size=17, color=BURGUNDY,
                 bold=True, align=PP_ALIGN.CENTER)
    elif kind == "disparities":
        items = [("POPULATION", "scale · distribution · mobility"), ("ECONOMY", "income · production · trade"),
                 ("HUMAN DEVELOPMENT", "poverty · health · education")]
        for i, (head, body) in enumerate(items):
            y = 1.58 + i * 1.38
            box(slide, 1.15, y, 11.0, 1.0, fill=PAPER, line=[FOREST, GOLD, BURGUNDY][i])
            text_box(slide, 1.48, y + 0.18, 2.55, 0.35, head, size=13, bold=True,
                     color=[FOREST, GOLD, BURGUNDY][i])
            text_box(slide, 4.12, y + 0.13, 7.45, 0.48, body, size=20, bold=True, font=TITLE_FONT)
        box(slide, 2.25, 5.92, 8.85, 0.5, "Keep the comparison · replace every inherited number before F4",
            fill=PALE_GOLD, line=GOLD, size=15, color=BURGUNDY, bold=True)
    elif kind == "institutions":
        box(slide, 0.9, 1.55, 5.4, 3.95, fill=PALE_GREEN, line=FOREST)
        text_box(slide, 1.25, 1.9, 4.7, 0.4, "FORMAL AMBITION", size=15, color=FOREST, bold=True,
                 align=PP_ALIGN.CENTER)
        text_box(slide, 1.42, 2.72, 4.35, 1.6, "SAARC\nSAFTA\ntrade · transit · cooperation",
                 size=22, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 7.0, 1.55, 5.4, 3.95, fill=PALE_RED, line=BURGUNDY)
        text_box(slide, 7.35, 1.9, 4.7, 0.4, "PRACTICAL LIMITS", size=15, color=BURGUNDY, bold=True,
                 align=PP_ALIGN.CENTER)
        text_box(slide, 7.52, 2.72, 4.35, 1.6, "implementation gaps\nbarriers to exchange\npolitical conflict",
                 size=22, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        box(slide, 2.0, 5.78, 9.35, 0.58, "Institutions prove organised intent—not necessarily regional capacity.",
            fill=PAPER, line=GOLD, size=16, color=INK, bold=True)
    elif kind == "matrix":
        rows = [("GEOGRAPHICAL FRAME", "STRONGER", FOREST), ("FUNCTIONAL INTERDEPENDENCE", "MATERIAL", FOREST),
                ("REGIONAL IDENTITY", "CONTESTED", GOLD), ("INSTITUTIONAL DEPTH", "LIMITED", BURGUNDY),
                ("COLLECTIVE ACTORNESS", "PROVISIONAL", BURGUNDY)]
        for i, (label, status, color) in enumerate(rows):
            y = 1.46 + i * 0.94
            box(slide, 1.12, y, 8.0, 0.68, label, fill=PAPER, line=LINE, size=15, color=INK,
                bold=True, align=PP_ALIGN.LEFT)
            box(slide, 9.42, y, 2.75, 0.68, status, fill=PALE_GREEN if color == FOREST else
                PALE_GOLD if color == GOLD else PALE_RED, line=color, size=14, color=color, bold=True)
        text_box(slide, 2.1, 6.28, 9.05, 0.32, "Analytical judgement—not a numerical score.", size=13,
                 color=MUTED, align=PP_ALIGN.CENTER)
    elif kind == "evidence":
        columns = [
            ("SUPPORTED LOCALLY", "region concept\nriver and functional cases", PALE_GREEN, FOREST),
            ("PARTIAL / DATED", "trade and institutions\ndevelopment comparisons", PALE_GOLD, GOLD),
            ("SOURCE NEEDED", "historical synthesis\nidentity · current actorness", PALE_RED, BURGUNDY),
        ]
        for i, (head, body, fill, color) in enumerate(columns):
            x = 0.8 + i * 4.18
            box(slide, x, 1.62, 3.75, 3.8, fill=fill, line=color)
            text_box(slide, x + 0.25, 1.98, 3.25, 0.42, head, size=13, color=color, bold=True,
                     align=PP_ALIGN.CENTER)
            text_box(slide, x + 0.35, 2.92, 3.05, 1.45, body, size=20, bold=True, font=TITLE_FONT,
                     align=PP_ALIGN.CENTER)
        box(slide, 1.75, 5.82, 9.85, 0.55, "Accepted master identity ≠ accepted academic claim",
            fill=PAPER, line=BURGUNDY, size=17, color=BURGUNDY, bold=True)
    elif kind == "conclusion":
        statements = [
            "South Asia is a meaningful field of spatial and historical enquiry.",
            "Its connections are real, but they are unequal and contested.",
            "Its institutional depth and collective actorness require cautious judgement.",
        ]
        for i, statement in enumerate(statements):
            y = 1.48 + i * 1.15
            text_box(slide, 1.05, y, 0.55, 0.55, str(i + 1), size=24, color=BURGUNDY,
                     bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
            text_box(slide, 1.85, y - 0.02, 10.25, 0.72, statement, size=21,
                     bold=True, font=TITLE_FONT)
        box(slide, 1.25, 5.18, 10.75, 1.08,
            "DISCUSS  ·  Which dimension makes South Asia most convincingly a region—and what evidence would change your answer?",
            fill=PALE_GOLD, line=GOLD, size=17, color=BURGUNDY, bold=True)
    elif kind == "references":
        box(slide, 0.85, 1.48, 4.05, 4.82, fill=PALE_GREEN, line=FOREST)
        text_box(slide, 1.18, 1.85, 3.4, 0.42, "FRIDAY BRIDGE", size=14, color=FOREST, bold=True,
                 align=PP_ALIGN.CENTER)
        text_box(slide, 1.28, 2.65, 3.2, 2.25,
                 "THE HIMALAYAS\n\nboundary\ncorridor\necological system\npolitical space",
                 size=21, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        text_box(slide, 5.35, 1.5, 6.9, 0.4, "SELECTED REFERENCES", size=14, color=BURGUNDY, bold=True)
        refs = (
            "Schmitt-Egner, P. (2002). “The Concept of ‘Region’.”\n\n"
            "Bandyopadhyay, S. et al., eds. (2017). Regional Cooperation in South Asia.\n\n"
            "Asian Development Bank (2009). Study on Intraregional Trade and Investment in South Asia.\n\n"
            "Ahmed, S., Kelegama, S. & Ghani, E., eds. (2010). Promoting Economic Cooperation in South Asia.\n\n"
            "Historical Lecture 1A pilot master (preserved F0 teaching artefact)."
        )
        text_box(slide, 5.35, 2.1, 6.75, 3.95, refs, size=14.2, font=BODY_FONT, valign=MSO_ANCHOR.TOP)


def build():
    DECK_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "South Asia as a Region — Lecture 1A v0.5"
    prs.core_properties.subject = "Synoptic companion to the accepted 65-slide F0 pilot master"
    prs.core_properties.author = "IS529N technical operator"
    prs.core_properties.comments = "Not approved for classroom use"
    for number, item in enumerate(SLIDES, 1):
        if item["kind"] == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            fill = slide.background.fill
            fill.solid(); fill.fore_color.rgb = IVORY
            render(slide, item)
        else:
            slide = base_slide(prs, number, item["title"], item["subtitle"], item["source"])
            render(slide, item)
        add_notes(slide, item)
    prs.save(PPTX)

    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    metadata = {
        "course_code": "IS529N",
        "lecture_identifier": "IS529N-L01-A",
        "version": "v0.5",
        "status": "SYNOPTIC_COMPANION",
        "review_status": "REVIEW_REQUIRED",
        "classroom_use_status": "NOT_FOR_CLASSROOM_USE",
        "fidelity_status": "F1",
        "canonical_master": "IS529N-RES-0001",
        "canonical_source_files_modified": False,
        "slide_count": len(SLIDES),
        "generated_at": generated,
        "pptx_sha256": sha256(PPTX),
        "academic_role": "SYNOPTIC_VIEW_SUPPORTING_TEACHER_INCUBATION_OF_MASTER",
        "design_direction": "simple classic restrained",
        "removed_or_excluded": [
            "unsupported ten-regions claim",
            "placeholder <number>",
            "unverified inherited numerical charts",
            "duplicate theory treatment",
            "planning-only slide prose",
        ],
        "slides": [
            {
                "sequence": i,
                "title": item["title"],
                "purpose": item["purpose"],
                "source": item["source"],
                "verification_status": item["verification"],
            }
            for i, item in enumerate(SLIDES, 1)
        ],
    }
    METADATA.write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    brief = [
        "# Lecture 1A v0.5 teaching brief",
        "",
        "- Status: `REVIEW_REQUIRED`",
        "- Fidelity: `F1` — structured synopsis, not factually approved",
        "- Academic role: synoptic companion to the 65-slide master",
        "- Classroom use: `NOT_FOR_CLASSROOM_USE`",
        "- Central question: **In what senses is South Asia a region?**",
        "- Working conclusion: South Asia is a layered and contested regional formation.",
        "",
        "## Sequence",
        "",
    ]
    for i, item in enumerate(SLIDES, 1):
        brief.extend([
            f"### {i}. {item['title']}",
            "",
            f"- Purpose: {item['purpose']}",
            f"- Speaker cue: {item['cue']}",
            f"- Evidence: {item['source']}",
            f"- Verification: `{item['verification']}`",
            "",
        ])
    brief.extend([
        "## Review gate",
        "",
        "This synopsis supports teacher incubation of the 65-slide master. It does not",
        "replace, abbreviate or become the revised master. Any classroom use still",
        "requires separate instructor review and factual verification.",
        "",
    ])
    BRIEF.write_text("\n".join(brief), encoding="utf-8")

    findings = [
        {
            "severity": "WARNING",
            "code": "FACTUAL_REVIEW_REQUIRED",
            "scope": "DECK",
            "message": "Historical, political and institutional claims remain explicitly qualified pending review.",
        },
        {
            "severity": "WARNING",
            "code": "CURRENT_DATA_REQUIRED",
            "scope": "SLIDE_10",
            "message": "No inherited numerical indicator is authorised for classroom use; current data must be commissioned separately.",
        },
        {
            "severity": "WARNING",
            "code": "CLASSROOM_APPROVAL_REQUIRED",
            "scope": "DECK",
            "message": "The synoptic companion is F1 and not approved for classroom use.",
        },
    ]
    report = {
        "status": "REVIEW_REQUIRED",
        "generated_at": generated,
        "slide_count": len(SLIDES),
        "error_count": 0,
        "warning_count": len(findings),
        "findings": findings,
    }
    VALIDATION.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(PPTX)
    print(METADATA)
    print(BRIEF)
    print(VALIDATION)


if __name__ == "__main__":
    build()
