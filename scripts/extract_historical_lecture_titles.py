#!/usr/bin/env python3
"""Extract title-slide evidence from authorised historical lecture decks."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CENSUS = Path("reports/source_census.json")
OUTPUT_JSON = Path("reports/historical_lecture_title_extraction.json")
OUTPUT_MD = Path("reports/historical_lecture_title_extraction.md")
HISTORICAL_LABEL = "<HISTORICAL_RESOURCE_REPOSITORY>"
DECK_FUNCTIONS = {
    "principal instructor lecture deck or drawing",
    "previous course-run version or backup",
}
COURSE_IDENTITY = "Economic and Political Geography of South Asia"


def atomic_write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def clean_text(value: str) -> str:
    value = value.replace("\x00", " ").replace("<number>", " ")
    return re.sub(r"\s+", " ", value).strip()


def pdf_pages(path: Path) -> list[str]:
    pages = []
    for number in range(1, 4):
        result = subprocess.run(
            ["pdftotext", "-f", str(number), "-l", str(number), "-layout", str(path), "-"],
            capture_output=True, text=True, timeout=30,
        )
        pages.append(clean_text(result.stdout))
    return pages


def odp_pages(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    pages = []
    for page in root.iter():
        if not page.tag.endswith("}page"):
            continue
        chunks = []
        for node in page.iter():
            if node.tag.endswith(("}p", "}h")):
                text = clean_text(" ".join(node.itertext()))
                if text and text not in chunks:
                    chunks.append(text)
        pages.append(clean_text(" ".join(chunks)))
        if len(pages) == 3:
            break
    return pages


def pptx_pages(path: Path) -> list[str]:
    from pptx import Presentation

    presentation = Presentation(path)
    pages = []
    for slide in list(presentation.slides)[:3]:
        chunks = [shape.text for shape in slide.shapes if hasattr(shape, "text_frame") and shape.text.strip()]
        pages.append(clean_text(" ".join(chunks)))
    return pages


def infer_filename_title(path: Path) -> str:
    title = re.sub(r"(?i)\b(?:lecture|old|backup)\s*\d*[a-d]?\b", " ", path.stem)
    title = re.sub(r"[_-]+", " ", title)
    return clean_text(title).title() or "Unresolved title"


def displayed_lecture_number(text: str) -> int | None:
    match = re.search(r"(?i)\blecture\s*(\d{1,2})", text)
    return int(match.group(1)) if match else None


def source_classification(text: str, relative: str) -> str:
    if re.search(r"(?i)(?:\bpart\s*[-–—]?\s*c\b|\blecture\s*\d+\s*c\b)", text) or re.search(r"(?i)lecture\d+c", Path(relative).stem):
        return "SUPPLEMENTARY_PART_C"
    if re.search(r"(?i)\bpart\s*[-–—]?\s*a\b", text) or re.search(r"(?i)lecture\d+a", Path(relative).stem):
        return "PART_A"
    if re.search(r"(?i)\bpart\s*[-–—]?\s*b\b", text) or re.search(r"(?i)lecture\d+b", Path(relative).stem):
        return "PART_B"
    return "WEEKLY_OR_UNCERTAIN"


def associated_lecture(relative: str, displayed: int | None) -> tuple[int | None, bool]:
    top = relative.split("/", 1)[0]
    folder_match = re.search(r"(?i)LECTURE(\d{1,2})", top)
    folder = int(folder_match.group(1)) if folder_match else None
    if top.upper() == "LECTURE7_D" and displayed == 9:
        return 9, False
    mismatch = bool(folder and displayed and folder != displayed)
    return folder or displayed, mismatch


def extract_displayed_title(text: str, filename: Path) -> tuple[str, str]:
    if not text:
        return infer_filename_title(filename), "TITLE_INFERRED_FROM_FILENAME"
    match = re.search(r"(?i)\blecture\s*\d{1,2}\s*[a-d]?\b", text)
    if not match:
        candidate = text.split(" Ambrish", 1)[0]
        candidate = candidate.replace(COURSE_IDENTITY, " ")
        candidate = clean_text(candidate)
        return (candidate or infer_filename_title(filename),
                "TITLE_EXTRACTED_FROM_SLIDE" if candidate else "TITLE_INFERRED_FROM_FILENAME")
    before = clean_text(text[:match.start()].replace(COURSE_IDENTITY, " "))
    after = text[match.end():].split("Ambrish", 1)[0]
    after = re.sub(r"(?i)^\s*[-–—]?\s*(?:part\s*[-–—]?\s*[a-c])\b", " ", after)
    after = clean_text(after)
    if before.casefold() == "political geography of south asia".casefold() and after:
        candidate = after
    elif before and after and after.casefold() not in before.casefold():
        candidate = f"{before}: {after}"
    else:
        candidate = before or after
    candidate = re.sub(r"\s+([,:;])", r"\1", clean_text(candidate))
    return (candidate or infer_filename_title(filename),
            "TITLE_EXTRACTED_FROM_SLIDE" if candidate else "TITLE_INFERRED_FROM_FILENAME")


def parse_registry(path: Path) -> dict[str, str]:
    records = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        identifier, title = line.split("|", 1)
        records[identifier] = title
    return records


def population_records(root: Path, evidence: list[dict[str, object]]) -> list[dict[str, object]]:
    combined = {
        **parse_registry(root / "course/lecture_titles.txt"),
        **parse_registry(root / "course/lecture_part_titles.txt"),
    }
    population = []
    for identifier, title in sorted(combined.items()):
        if title == "Pending instructor confirmation":
            method = "UNRESOLVED_TITLE"
        else:
            matching = [record for record in evidence if record["target_identifier"] == identifier]
            if any(str(record["extracted_title"]).casefold() == title.casefold() for record in matching):
                method = "TITLE_EXTRACTED_FROM_SLIDE"
            elif identifier.endswith(tuple("AB")) and matching:
                method = "TITLE_RECONSTRUCTED_FROM_COURSE_SEQUENCE"
            elif len(identifier) == 3:
                part_evidence = [record for record in evidence if str(record["target_identifier"]).startswith(identifier) and len(str(record["target_identifier"])) == 4]
                method = "TITLE_RECONSTRUCTED_FROM_COURSE_SEQUENCE" if part_evidence else "TITLE_INFERRED_FROM_FILENAME"
            else:
                method = "TITLE_INFERRED_FROM_FILENAME"
        population.append({
            "identifier": identifier,
            "candidate_title": title,
            "population_method": method,
            "status": "HISTORICAL_CANDIDATE",
            "instructor_confirmation_required": True,
        })
    return population


def extract(root: Path, *, write: bool = True) -> dict[str, object]:
    census = json.loads((root / CENSUS).read_text(encoding="utf-8"))
    source_root = Path(census["scope"]["historical_repository"])
    selected = [
        record for record in census["records"]
        if record.get("probable_function") in DECK_FUNCTIONS
        and str(record.get("extension") or "").lower() in {".pdf", ".odp", ".pptx"}
        and Path(record["absolute_path"]).is_relative_to(source_root)
    ]
    evidence = []
    for record in selected:
        path = Path(record["absolute_path"])
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != record["content_sha256"]:
            raise ValueError(f"historical checksum mismatch: {record['relative_context']}")
        extension = path.suffix.lower()
        pages = pdf_pages(path) if extension == ".pdf" else odp_pages(path) if extension == ".odp" else pptx_pages(path)
        substantive = next((page for page in pages if page), "")
        title, method = extract_displayed_title(substantive, path)
        displayed = displayed_lecture_number(substantive)
        lecture_number, mismatch = associated_lecture(record["relative_context"], displayed)
        classification = source_classification(substantive, record["relative_context"])
        if classification == "PART_A":
            target = f"L{lecture_number:02d}A" if lecture_number else "UNRESOLVED"
        elif classification == "PART_B":
            target = f"L{lecture_number:02d}B" if lecture_number else "UNRESOLVED"
        elif classification == "SUPPLEMENTARY_PART_C":
            target = f"L{lecture_number:02d}C" if lecture_number else "UNRESOLVED"
        else:
            target = f"L{lecture_number:02d}" if lecture_number else "UNRESOLVED"
        confidence = "HIGH" if method == "TITLE_EXTRACTED_FROM_SLIDE" and not mismatch else "MEDIUM" if method == "TITLE_EXTRACTED_FROM_SLIDE" else "LOW"
        evidence.append({
            "associated_lecture_id": f"IS529N-L{lecture_number:02d}" if lecture_number else None,
            "target_identifier": target,
            "historical_classification": classification,
            "source_filename": path.name,
            "source_locator": f"{HISTORICAL_LABEL}/{record['relative_context']}",
            "source_sha256": digest,
            "extracted_title": title,
            "extraction_method": method,
            "extraction_confidence": confidence,
            "displayed_lecture_number": displayed,
            "repository_sequence_mismatch": mismatch,
            "instructor_confirmation_required": True,
        })
    population = population_records(root, evidence)
    counts = {}
    for record in evidence:
        counts[record["extraction_method"]] = counts.get(record["extraction_method"], 0) + 1
    payload = {
        "report_type": "READ_ONLY_HISTORICAL_LECTURE_TITLE_EXTRACTION",
        "historical_source_files_modified": False,
        "bulk_ocr_performed": False,
        "record_count": len(evidence),
        "method_counts": counts,
        "records": evidence,
        "registry_population": population,
    }
    lines = [
        "# Historical Lecture Title Extraction", "",
        "Read-only extraction from authorised title slides or first substantive slides. No OCR was performed. All results remain historical candidates pending instructor confirmation.", "",
        f"- Deck records inspected: {len(evidence)}",
        f"- Titles extracted from slide text: {counts.get('TITLE_EXTRACTED_FROM_SLIDE', 0)}",
        f"- Titles inferred from filename: {counts.get('TITLE_INFERRED_FROM_FILENAME', 0)}", "",
        "| Target | Classification | Extracted title | Method | Confidence | Source | SHA-256 | Confirmation |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for record in evidence:
        title = str(record["extracted_title"]).replace("|", "\\|")
        lines.append(
            f"| {record['target_identifier']} | {record['historical_classification']} | {title} | "
            f"{record['extraction_method']} | {record['extraction_confidence']} | {record['source_filename']} | "
            f"`{record['source_sha256']}` | Required |"
        )
    lines.extend(["", "## Registry population", "", "| Identifier | Candidate title | Basis | Status |", "|---|---|---|---|"])
    for record in population:
        lines.append(f"| {record['identifier']} | {record['candidate_title']} | {record['population_method']} | HISTORICAL_CANDIDATE |")
    if write:
        atomic_write(root / OUTPUT_JSON, json.dumps(payload, indent=2, sort_keys=True) + "\n")
        atomic_write(root / OUTPUT_MD, "\n".join(lines) + "\n")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--dry-run", action="store_true")
    arguments = parser.parse_args()
    result = extract(arguments.project_root.resolve(), write=not arguments.dry_run)
    print(f"historical decks inspected: {result['record_count']}")
    print(json.dumps(result["method_counts"], sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
